"""Excel-Übersicht aller Unternehmen: was ist erledigt, was muss ich noch selbst prüfen?

Blatt „Firmen“: eine Zeile pro Unternehmen mit Kontaktdaten, Status jedes Schritts, Score,
Grund für Übersprungen/Fehler, Design-Urteil und der Spalte „Manuell prüfen“.
Blatt „Befunde“: alle Befunde der jeweils letzten Analyse.

„Manuell prüfen = Ja“ heißt: Es gibt keinen fertigen Entwurf (übersprungen, ausgeschlossen
oder fehlgeschlagen). Die Spalte „Grund“ sagt, warum.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from modules.database import Database, Step
from modules.design_review import load_review

STATUS_DE = {"done": "erledigt", "skipped": "übersprungen", "failed": "Fehler", "running": "läuft",
             "pending": "offen", None: "–"}
STEP_DE = {Step.CRAWL: "Crawl", Step.ANALYZE: "Analyse", Step.SCORE: "Bewertung", Step.DRAFT: "Entwurf"}
KIND_DE = {"own": "eigene", "directory": "nur Verzeichnis/Social", "none": "keine"}
YELLOW = PatternFill("solid", fgColor="FFF2CC")
GREEN = PatternFill("solid", fgColor="E2EFDA")
HEADER = PatternFill("solid", fgColor="D9D9D9")

COLUMNS = [
    ("ID", 6), ("Firma", 34), ("Branche", 18), ("Ort", 16), ("Telefon", 17), ("E-Mail", 30), ("Website", 34),
    ("Website-Art", 14), ("Manuell prüfen", 10), ("Grund", 50), ("Crawl", 12), ("Analyse", 12), ("Entwurf", 12),
    ("Score", 7), ("Wichtigste Befunde", 60), ("Design-Urteil", 18), ("Design: schlecht", 50),
    ("Letzter Kontakt", 22), ("Entwurf-Datei", 30),
]


def _steps(db: Database, company_id: int) -> dict[str, dict]:
    rows = db.conn.execute("SELECT step, status, last_error FROM job_steps WHERE company_id = ?", (company_id,))
    return {r["step"]: {"status": r["status"], "error": r["last_error"]} for r in rows}


def _reason(steps: dict[str, dict]) -> str:
    """Erster Schritt, der übersprungen wurde oder scheiterte – mit Grund (die eigentliche Ursache)."""
    for step in (Step.CRAWL, Step.ANALYZE, Step.SCORE, Step.DRAFT):
        info = steps.get(step.value)
        if info is None:
            return f"{STEP_DE[step]}: noch nicht gelaufen"
        if info["status"] in ("skipped", "failed"):
            return f"{STEP_DE[step]} {STATUS_DE[info['status']]}: {info['error'] or '–'}"
    return ""


def collect_rows(db: Database, data_dir: Path) -> tuple[list[list], list[list]]:
    firms, findings = [], []
    companies = db.conn.execute("SELECT * FROM companies ORDER BY id").fetchall()
    for c in companies:
        cid = c["id"]
        steps = _steps(db, cid)
        score_row = db.conn.execute("SELECT * FROM lead_scores WHERE company_id = ? ORDER BY id DESC LIMIT 1",
                                    (cid,)).fetchone()
        draft = db.conn.execute("SELECT recipient, eml_path FROM email_drafts WHERE company_id = ? ORDER BY id DESC "
                                "LIMIT 1", (cid,)).fetchone()
        contact = db.conn.execute("SELECT occurred_at, channel, outcome FROM contact_events WHERE company_id = ? "
                                  "ORDER BY occurred_at DESC LIMIT 1", (cid,)).fetchone()
        emails = [r["value"] for r in db.get_contacts(cid, "EMAIL")]
        email = (draft["recipient"] if draft else None) or c["email"] or (emails[0] if emails else None)
        draft_done = (steps.get(Step.DRAFT.value) or {}).get("status") == "done"

        report_id = db.latest_report_id(cid)
        report = db.get_report(report_id) if report_id else None
        top = []
        for f in (report.findings if report else []):
            findings.append([cid, c["name"], f.category.value, f.severity.value, f.description, f.evidence,
                             f.source_url, f.evidence_type.value, f.confidence])
            if f.severity.value in ("HIGH", "MEDIUM"):
                top.append(f.description)

        design = load_review(data_dir, c["domain"])
        firms.append([
            cid, c["name"], c["niche_branch"] or c["category"], c["city"], c["phone"], email, c["website"],
            KIND_DE.get(c["website_kind"], c["website_kind"] or "–"),
            "Nein" if draft_done else "Ja", "" if draft_done else _reason(steps),
            *(STATUS_DE.get((steps.get(s.value) or {}).get("status"), "–") for s in (Step.CRAWL, Step.ANALYZE, Step.DRAFT)),
            None if score_row is None or score_row["disqualified"] else round(score_row["score"]),
            " | ".join(top[:3]),
            design.summary if design else "",
            " | ".join(design.bad[:4]) if design else "",
            f"{contact['occurred_at'][:10]} {contact['channel']} → {contact['outcome']}" if contact else "",
            draft["eml_path"] if draft and draft["eml_path"] else "",
        ])
    # Erst alles, was ich mir selbst ansehen muss; darin die mit Website zuerst
    firms.sort(key=lambda r: (r[8] != "Ja", r[7] != "eigene", -(r[13] or 0), r[0]))
    return firms, findings


def _sheet(ws, header: list[tuple[str, int]], rows: list[list]) -> None:
    ws.append([h for h, _ in header])
    for i, (_, width) in enumerate(header, start=1):
        cell = ws.cell(row=1, column=i)
        cell.font, cell.fill = Font(bold=True), HEADER
        ws.column_dimensions[get_column_letter(i)].width = width
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)


def export_excel(db: Database, path: Path, data_dir: Path) -> tuple[Path, int, int]:
    """Schreibt die Excel-Datei. Liefert (Pfad, Anzahl Firmen, davon manuell zu prüfen)."""
    firms, findings = collect_rows(db, data_dir)
    wb = Workbook()
    ws = wb.active
    ws.title = "Firmen"
    _sheet(ws, COLUMNS, firms)
    manual_col = [h for h, _ in COLUMNS].index("Manuell prüfen")
    for row in ws.iter_rows(min_row=2):
        fill = YELLOW if row[manual_col].value == "Ja" else GREEN
        for cell in row:
            cell.fill = fill
    _sheet(wb.create_sheet("Befunde"), [("ID", 6), ("Firma", 34), ("Kategorie", 16), ("Schwere", 9),
                                        ("Beschreibung", 60), ("Beleg", 50), ("Fundstelle", 40), ("Belegart", 18),
                                        ("Sicherheit", 9)], findings)
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path, len(firms), sum(1 for r in firms if r[8] == "Ja")
