"""Gesamtablauf über alle Schritte mit lokalem Testserver und simulierter KI."""

import asyncio
import email

import pytest

from config import Settings
from core.niche import load_niche
from main import apply_overrides, main
from models.schemas import Lead
from modules.database import Database, Step, StepStatus
from modules.llm import Usage
from modules.pitch import Sender
from services.analysis_service import run_analysis
from services.crawl_service import run_crawls
from services.outreach_service import run_scoring_and_drafts


class EmptyLLM:
    model = "fake"

    def complete_json(self, system, user, max_tokens=1500):
        return {"findings": []}, Usage(10, 5, 0.0)


def make_settings(tmp_path, **kw):
    return Settings(_env_file=None, data_dir=tmp_path, logs_dir=tmp_path / "logs", database_path=tmp_path / "t.db",
                    crawl_max_pages=4, crawl_timeout_seconds=2, crawl_retries=0, crawl_max_seconds_per_site=30, **kw)


def add_company(db, name, url, kind="own"):
    cid = db.upsert_company(Lead(name=name, source="test", source_ref=name, niche_fit=1.0,
                                 niche_branch="Schreiner/Tischler")).company_id
    # lokale Test-URLs lehnt Lead (zu Recht) ab – deshalb direkt setzen
    db.conn.execute("UPDATE companies SET website = ?, website_kind = ?, domain = ? WHERE id = ?",
                    (url, kind, f"test-{cid}.local" if url else None, cid))
    return cid


@pytest.fixture
def env(tmp_path, site_server):
    settings = make_settings(tmp_path)
    db = Database(settings.database_path)
    good = add_company(db, "Schreinerei Beispiel", site_server + "/")
    blocked = add_company(db, "Gesperrt GmbH", site_server + "/forbidden")
    jsonly = add_company(db, "Nur JS", site_server + "/js-only.html")
    directory = add_company(db, "Profil Maler", "https://www.firmen.tv/x", kind="directory")
    yield db, settings, {"good": good, "blocked": blocked, "js": jsonly, "dir": directory}
    db.close()


def run_all(db, settings):
    crawl = asyncio.run(run_crawls(db, settings, use_browser=False, allow_local=True))
    analysis = asyncio.run(run_analysis(db, settings, EmptyLLM(), network_probes=False))
    scored, drafted = run_scoring_and_drafts(db, settings, load_niche(), Sender())
    return crawl, analysis, scored, drafted


def test_full_pipeline(env):
    db, settings, ids = env
    crawl, analysis, scored, drafted = run_all(db, settings)

    assert db.step_status(ids["dir"], Step.CRAWL) is StepStatus.SKIPPED
    assert db.step_status(ids["blocked"], Step.CRAWL) is StepStatus.DONE          # dauerhaft: nicht wiederholen
    assert db.latest_crawl(ids["blocked"])[1].status.value == "BLOCKED"

    # Kernregel: unbrauchbare Crawls werden nie als Analyse gespeichert
    for key in ("blocked", "js"):
        assert db.step_status(ids[key], Step.ANALYZE) is StepStatus.SKIPPED
        assert db.latest_report_id(ids[key]) is None
    assert db.step_status(ids["good"], Step.ANALYZE) is StepStatus.DONE

    # Nur der gute Lead bekommt einen Entwurf
    drafts = db.list_drafts()
    assert [d["company_id"] for d in drafts] == [ids["good"]]
    row = drafts[0]
    assert row["recipient"] == "info@schreinerei-beispiel.de" and row["status"] == "draft"
    msg = email.message_from_bytes((settings.data_dir / row["eml_path"]).read_bytes())
    body = msg.get_payload(decode=True).decode()
    assert "keine weiteren Nachrichten" in body and "fieberki.de" in body
    assert any("ausgeschlossen" in d and "Profil Maler" in d for d in scored.details)


def test_rerun_does_not_repeat_finished_steps(env):
    db, settings, ids = env
    run_all(db, settings)
    crawl, analysis, scored, drafted = run_all(db, settings)
    assert (crawl.done, analysis.done, scored.done, drafted.done) == (0, 0, 0, 0)
    assert len(db.list_drafts()) == 1
    runs = db.conn.execute("SELECT COUNT(*) FROM crawl_runs WHERE company_id = ?", (ids["good"],)).fetchone()[0]
    assert runs == 1


def test_retry_single_lead(env):
    db, settings, ids = env
    run_all(db, settings)
    db.reset_step(ids["good"], Step.CRAWL)
    crawl = asyncio.run(run_crawls(db, settings, company_ids=[ids["good"]], use_browser=False, allow_local=True))
    assert crawl.done == 1


def test_one_broken_company_does_not_stop_run(env, monkeypatch):
    db, settings, ids = env
    import services.crawl_service as cs
    original = cs.SiteCrawler.crawl_site

    async def broken(self, url):
        if "js-only" in url:
            raise RuntimeError("Absturz")
        return await original(self, url)
    monkeypatch.setattr(cs.SiteCrawler, "crawl_site", broken)
    report = asyncio.run(run_crawls(db, settings, use_browser=False, allow_local=True))
    assert report.failed == 1 and report.done == 2
    assert db.errors_for(ids["js"])[0]["error_type"] == "RuntimeError"
    assert db.step_status(ids["js"], Step.CRAWL) is StepStatus.FAILED


# --- Kommandozeile ---

def test_cli_status_leads_drafts(env, capsys):
    db, settings, ids = env
    run_all(db, settings)
    for cmd in (["status"], ["leads"], ["drafts"], ["contacts", "list"]):
        assert main(cmd, settings=settings) == 0
    out = capsys.readouterr().out
    assert "Unternehmen: 4" in out and "Schreinerei Beispiel" in out and "entwurf_" in out


def test_cli_send_refused_by_default(env, capsys):
    db, settings, ids = env
    run_all(db, settings)
    draft_id = db.list_drafts()[0]["id"]
    assert main(["send", "--draft", str(draft_id)], settings=settings) == 2
    assert "deaktiviert" in capsys.readouterr().out
    assert db.get_draft(draft_id)["status"] == "draft"


def test_cli_contact_and_block(env, capsys):
    db, settings, ids = env
    assert main(["contact", "--lead", str(ids["good"]), "--kanal", "telefon", "--ergebnis", "kein_interesse"],
                settings=settings) == 0
    assert main(["block", "--domain", "spam.de", "--grund", "test"], settings=settings) == 0
    db2 = Database(settings.database_path)
    assert db2.get_company(ids["good"])["do_not_contact"] == 1 and db2.is_blocked(domain="spam.de")
    db2.close()


def test_apply_overrides():
    niche = apply_overrides(load_niche(), "Burgau", "Maler")
    assert niche.region.orte == ["Burgau"] and [b.name for b in niche.branche] == ["Maler/Stuckateur"]
    with pytest.raises(SystemExit):
        apply_overrides(load_niche(), None, "Bäcker")


# --- Excel-Übersicht und Oberfläche ---

def test_excel_shows_what_needs_manual_review(env, tmp_path):
    from openpyxl import load_workbook

    from modules.excel_export import export_excel
    db, settings, ids = env
    run_all(db, settings)
    path, total, manual = export_excel(db, tmp_path / "firmen.xlsx", settings.data_dir)
    assert (total, manual) == (4, 3)
    ws = load_workbook(path)["Firmen"]
    header = [c.value for c in ws[1]]
    rows = {r[header.index("ID")]: dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)}
    assert rows[ids["good"]]["Manuell prüfen"] == "Nein" and rows[ids["good"]]["Entwurf"] == "erledigt"
    assert rows[ids["dir"]]["Manuell prüfen"] == "Ja" and "Verzeichnis" in rows[ids["dir"]]["Grund"]
    assert rows[ids["blocked"]]["Manuell prüfen"] == "Ja" and "BLOCKED" in rows[ids["blocked"]]["Grund"]
    first = next(ws.iter_rows(min_row=2, max_row=2, values_only=True))
    assert first[header.index("Manuell prüfen")] == "Ja"                    # Manuelles steht oben
    assert load_workbook(path)["Befunde"].max_row >= 2


def test_main_export_command(env, capsys):
    db, settings, _ = env
    assert main(["export"], settings=settings, niche=load_niche()) == 0
    assert (settings.data_dir / "firmen.xlsx").exists() and "manuell prüfen" in capsys.readouterr().out


def test_ui_check_report_and_downloads(env, site_server):
    import threading
    from http.server import ThreadingHTTPServer

    import httpx

    from modules.browser import BrowserLimits
    from modules.design_review import review_site
    from modules import web_ui
    db, settings, ids = env
    run_all(db, settings)

    async def local_review(url, data_dir, llm=None):          # Testserver läuft auf 127.0.0.1
        return await review_site(url, data_dir, llm, limits=BrowserLimits(navigation_timeout_seconds=5,
                                                                          networkidle_timeout_seconds=1,
                                                                          allow_local=True))
    original = web_ui.review_site
    web_ui.review_site = local_review
    server = ThreadingHTTPServer(("127.0.0.1", 0), web_ui.make_handler(settings, lambda: None))
    threading.Thread(target=server.serve_forever, daemon=True).start()
    base = f"http://127.0.0.1:{server.server_address[1]}"
    try:
        with httpx.Client(base_url=base, timeout=60) as http:
            index = http.get("/").text
            assert "Schreinerei Beispiel" in index and "manuell prüfen" in index
            r = http.post("/pruefen", data={"url": site_server + "/design-old.html"})
            assert r.status_code == 303 and r.headers["location"] == "/bericht?d=127.0.0.1"
            report = http.get("/bericht?d=127.0.0.1").text
            assert "sehr veraltet" in report and "Tabellen" in report
            shot = http.get("/datei?d=127.0.0.1&f=desktop.png&download=1")
            assert shot.content[:4] == b"\x89PNG" and "attachment" in shot.headers["content-disposition"]
            assert http.get("/datei?d=127.0.0.1&f=../../t.db").status_code == 404
            assert http.get("/excel").content[:2] == b"PK"
    finally:
        web_ui.review_site = original
        server.shutdown()
        server.server_close()
